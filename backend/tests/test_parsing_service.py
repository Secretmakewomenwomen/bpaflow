from io import BytesIO
from types import SimpleNamespace

from docx import Document

from app.services.parsing_service import ParsingService


class ParsingServiceHarness(ParsingService):
    def _extract_pdf_page_texts(self, content: bytes) -> list[str]:
        assert content == b"pdf-bytes"
        return ["", "第二页原生文本"]

    def _ocr_pdf_page_texts(self, content: bytes, page_indexes: list[int]) -> dict[int, str]:
        assert content == b"pdf-bytes"
        assert page_indexes == [0]
        return {0: "第一页OCR文本"}


def test_parse_pdf_uses_ocr_for_pages_without_native_text() -> None:
    service = ParsingServiceHarness(SimpleNamespace(ocr_language="chi_sim+eng"))

    document = service.parse(
        filename="diagram.pdf",
        file_ext="pdf",
        mime_type="application/pdf",
        content=b"pdf-bytes",
    )

    assert len(document.segments) == 2
    assert document.segments[0].text == "第一页OCR文本"
    assert document.segments[0].source_type == "pdf_ocr"
    assert document.segments[0].page_start == 1
    assert document.segments[1].text == "第二页原生文本"
    assert document.segments[1].source_type == "pdf_text"


def test_parse_docx_extracts_text_from_table_rows() -> None:
    document = Document()
    table = document.add_table(rows=2, cols=3)
    table.cell(0, 0).text = "流程名称"
    table.cell(0, 1).text = "3.0 洞察到营销策略"
    table.cell(0, 2).text = "L1"
    table.cell(1, 0).text = "流程定义"
    table.cell(1, 1).text = "定义市场洞察流程"
    table.cell(1, 2).text = "指导营销策略制定"
    buffer = BytesIO()
    document.save(buffer)

    service = ParsingService(SimpleNamespace(ocr_language="chi_sim+eng"))
    parsed = service.parse(
        filename="table-only.docx",
        file_ext="docx",
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        content=buffer.getvalue(),
    )

    assert [segment.text for segment in parsed.segments] == [
        "流程名称 | 3.0 洞察到营销策略 | L1",
        "流程定义 | 定义市场洞察流程 | 指导营销策略制定",
    ]
